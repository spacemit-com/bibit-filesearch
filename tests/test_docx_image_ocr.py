"""docx 内嵌图 OCR（rev 9）回归：与 PDF「逐页薄页判据」同构。

守四件事：
  ① 段落无文字 + 段内有图 → 对该图 OCR，文本进 text（+ images 留溯源）；
  ② 判据是**逐段**不是全篇：图与文字同段时不 OCR，无字段的图才 OCR；
  ③ 每文件 OCR 张数闸（DOCX_MAX_OCR_IMAGES）按「尝试张数」算，超出即 capped；
  ④ rev 必须递增——否则存量 docx 因 hash 不变永远停在 unsupported。

① 用**真实** OCR 调用（非 mock）；引擎不可用则 skip，不打桩，避免测了个假东西。
"""
from __future__ import annotations

import io
import shutil
import tempfile
from pathlib import Path

import pytest

from bibit_filesearch import parser as P


def _png_with_text(text: str = "极地冰川自动观测站 PG-88") -> bytes:
    """造一张带黑字的白底 PNG（供 OCR 真读）。"""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 220), "white")
    dr = ImageDraw.Draw(img)
    dr.text((30, 90), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _docx_image_only(blob: bytes) -> Path:
    """一个只有一张内嵌图（无文字）的 docx。"""
    import docx

    d = docx.Document()
    d.add_picture(io.BytesIO(blob))          # 新建一个空段落 + 图
    p = Path(tempfile.mkdtemp(prefix="docxrev9_")) / "img_only.docx"
    d.save(str(p))
    return p


def _docx_text_and_image_same_para(blob: bytes, caption: str) -> Path:
    """图与文字**同一段落**的 docx（用于验「有字的段不 OCR」）。"""
    import docx

    d = docx.Document()
    para = d.add_paragraph()
    para.add_run(caption)
    para.add_run().add_picture(io.BytesIO(blob))
    p = Path(tempfile.mkdtemp(prefix="docxrev9_")) / "mixed.docx"
    d.save(str(p))
    return p


def _ocr_available() -> bool:
    """OCR 引擎真可用才跑真读那一例（rapidocr 是 PP-OCR 的载体）。"""
    try:
        import rapidocr  # noqa: F401
        return True
    except Exception:                                # noqa: BLE001
        return False


def test_image_only_docx_gets_ocr_text():
    """① 纯图 docx：段落无文字 → OCR 该图，文本进 text（+ images 留溯源）。"""
    if not _ocr_available():
        pytest.skip("OCR 引擎不可用")
    p = _docx_image_only(_png_with_text())
    try:
        doc = P.parse_file(p)
        assert doc.file_type == "docx"
        assert doc.text.strip(), "内嵌图里的文字必须被 OCR 出来进 text；实际空"
        assert doc.images, "图生文关联（images）必须落，供溯源"
        assert doc.images[0]["ocr_text"].strip()
        assert doc.raw_meta["docx_scan"]["ocr_images"] >= 1
    finally:
        shutil.rmtree(p.parent, ignore_errors=True)


def test_text_paragraph_is_not_ocrd(monkeypatch):
    """② 判据逐段：图与文字同段 → 不 OCR；只有无字段的图才 OCR。"""
    blob = _png_with_text()
    # (a) 图与文字同段 → 该段不 OCR
    p_same = _docx_text_and_image_same_para(blob, "正文标题：观测记录")
    calls = {"n": 0}

    def _fake_ocr(img):
        calls["n"] += 1
        return ("不该被读到", 99.0)

    monkeypatch.setattr(P, "_ocr_image", _fake_ocr)
    try:
        doc = P.parse_file(p_same)
        assert "正文标题" in doc.text
        assert calls["n"] == 0, \
            "图与文字同段时不该 OCR（判据逐段，不是全篇）；实际尝试 %d 次" % calls["n"]
        # (b) 对照组：同图放空段 → 会被 OCR（证明 (a) 不是因为图取不到）
        calls["n"] = 0
        p_empty = _docx_image_only(blob)
        try:
            P.parse_file(p_empty)
            assert calls["n"] == 1, \
                "无字段的图必须被 OCR（对照组）；实际尝试 %d 次" % calls["n"]
        finally:
            shutil.rmtree(p_empty.parent, ignore_errors=True)
    finally:
        shutil.rmtree(p_same.parent, ignore_errors=True)


def test_ocr_image_cap_counts_attempts(monkeypatch):
    """③ 闸按「尝试张数」算：超出上限即停止尝试并标 capped。"""
    import docx

    d = docx.Document()
    blob = _png_with_text()
    for _ in range(3):
        d.add_picture(io.BytesIO(blob))
    p = Path(tempfile.mkdtemp(prefix="docxcap_")) / "cap.docx"
    d.save(str(p))

    calls = {"n": 0}

    def _fake_ocr(img):
        calls["n"] += 1
        return ("文本%d" % calls["n"], 99.0)

    monkeypatch.setattr(P, "_ocr_image", _fake_ocr)
    monkeypatch.setattr(P, "DOCX_MAX_OCR_IMAGES", 2)
    try:
        doc = P.parse_file(p)
        scan = doc.raw_meta["docx_scan"]
        assert calls["n"] == 2, "闸设 2 时只应尝试 2 张；实际 %d" % calls["n"]
        assert scan["ocr_attempts"] == 2
        assert scan["ocr_images_capped"] is True
    finally:
        shutil.rmtree(p.parent, ignore_errors=True)


def test_revision_bumped_to_9():
    """rev 必须递增——否则存量 docx 因 hash 不变永远停在 unsupported。"""
    assert P.PARSER_REVISION >= 9
